#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Automated Test Suite for UoPeople Coursework Unpacker & Assignment Template Population
Verifies:
1. Assignment template bundling in scripts/assets/template_assignment.docx
2. Population of Title, Department, Course, Instructor, and Due Date
3. Template copied ONLY into coursework/Assignment_Activity/ (never in Discussion_Forum)
4. Dynamic filename: week{N}_assignment_{initials}_template.docx
5. course_metadata.json parsing and integration
6. Companion daemon health and CORS endpoints
"""

import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace', line_buffering=True)
sys.stderr.reconfigure(encoding='utf-8', errors='replace')

import os
import re
import json
import time
import shutil
import tempfile
import unittest
import threading
import urllib.request
from pathlib import Path

# Add scripts directory to path
REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import process_active_courses as pac
import course_exporter_daemon as daemon


class TestAssignmentTemplatePopulation(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="uop_test_")
        self.bundled_template = SCRIPTS_DIR / "assets" / "template_assignment.docx"

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_bundled_template_exists(self):
        """Verify the base template is bundled into scripts/assets/template_assignment.docx."""
        self.assertTrue(self.bundled_template.is_file(), f"Missing bundled template at {self.bundled_template}")
        found = pac.find_assignment_template()
        self.assertIsNotNone(found)
        self.assertTrue(os.path.isfile(found))

    def test_populate_assignment_template_contents(self):
        """Verify population of Title, Department, Course, Instructor, and Due Date."""
        target_docx = os.path.join(self.temp_dir, "week4_assignment_myk_template.docx")
        title = "Unit 4 Assignment Activity"
        department = "Department of Computer Science, University of The People"
        course = "CS 2401: Software Engineering 1"
        instructor = "Christor Pancho"
        due_date = "October 15, 2026"

        success = pac.populate_assignment_template(
            base_template_path=str(self.bundled_template),
            target_path=target_docx,
            title=title,
            department=department,
            course=course,
            instructor=instructor,
            due_date=due_date
        )
        self.assertTrue(success, "Failed to populate assignment template")
        self.assertTrue(os.path.isfile(target_docx), "Target docx was not written")

        import docx
        doc = docx.Document(target_docx)
        self.assertEqual(doc.paragraphs[1].text.strip(), title)
        self.assertEqual(doc.paragraphs[3].text.strip(), department)
        self.assertEqual(doc.paragraphs[4].text.strip(), course)
        self.assertEqual(doc.paragraphs[5].text.strip(), instructor)
        self.assertEqual(doc.paragraphs[6].text.strip(), due_date)

    def test_course_folder_processing_and_isolation(self):
        """
        Verify that:
        - Template is copied ONLY into coursework/Assignment_Activity/ (never in Discussion_Forum)
        - Named week{N}_assignment_{initials}_template.docx
        - Uses metadata from course_metadata.json
        """
        course_dir = os.path.join(self.temp_dir, "CS2401-Software_Engineering_1")
        os.makedirs(course_dir, exist_ok=True)

        # Create sample course_metadata.json
        metadata = {
            "courseId": "123456",
            "courseCode": "CS 2401",
            "courseName": "Software Engineering 1",
            "department": "Computer Science",
            "departmentLine": "Department of Computer Science, University of The People",
            "courseLine": "CS 2401: Software Engineering 1",
            "instructor": "Dr. Christor Pancho",
            "student": {
                "name": "Mohamed Yahia Khidr",
                "initials": "myk"
            },
            "assignments": {
                "1": {
                    "unit": 1,
                    "title": "Unit 1 Assignment Activity",
                    "dueDate": "September 17, 2026",
                    "templateFileName": "week1_assignment_myk_template.docx"
                },
                "2": {
                    "unit": 2,
                    "title": "Unit 2 Assignment Activity",
                    "dueDate": "September 24, 2026",
                    "templateFileName": "week2_assignment_myk_template.docx"
                }
            }
        }
        with open(os.path.join(course_dir, "course_metadata.json"), "w", encoding="utf-8") as f:
            json.dump(metadata, f, indent=2)

        # Create Unit 1 and Unit 2 folders with discussion and assignment files
        for u in [1, 2]:
            u_folder = os.path.join(course_dir, f"0{u}_Unit_{u}")
            os.makedirs(u_folder, exist_ok=True)
            with open(os.path.join(u_folder, "03_Discussions.md"), "w", encoding="utf-8") as f:
                f.write(f"# Discussion Forum Unit {u}\nPlease respond to the discussion question in detail.")
            with open(os.path.join(u_folder, "04_Assignments.md"), "w", encoding="utf-8") as f:
                f.write(f"# Assignment Activity Unit {u}\nComplete the written assignment according to APA guidelines.")

        result = pac.process_course_folder(course_dir, self.temp_dir)
        self.assertTrue(result)

        for u in [1, 2]:
            u_folder = os.path.join(course_dir, f"0{u}_Unit_{u}")
            disc_dir = os.path.join(u_folder, "coursework", "Discussion_Forum")
            assign_dir = os.path.join(u_folder, "coursework", "Assignment_Activity")

            # 1. Discussion Forum exists but has NO template docx
            self.assertTrue(os.path.isdir(disc_dir), f"Discussion_Forum directory missing in Unit {u}")
            disc_files = os.listdir(disc_dir)
            self.assertEqual(len(disc_files), 0, f"Discussion_Forum MUST NOT contain template docx, found: {disc_files}")

            # 2. Assignment Activity exists and has populated template
            self.assertTrue(os.path.isdir(assign_dir), f"Assignment_Activity directory missing in Unit {u}")
            expected_filename = f"week{u}_assignment_myk_template.docx"
            expected_filepath = os.path.join(assign_dir, expected_filename)
            self.assertTrue(os.path.isfile(expected_filepath), f"Expected {expected_filename} in {assign_dir}")

            # 3. Check contents of populated template
            import docx
            doc = docx.Document(expected_filepath)
            self.assertEqual(doc.paragraphs[1].text.strip(), f"Unit {u} Assignment Activity")
            self.assertEqual(doc.paragraphs[3].text.strip(), "Department of Computer Science, University of The People")
            self.assertEqual(doc.paragraphs[4].text.strip(), "CS 2401: Software Engineering 1")
            self.assertEqual(doc.paragraphs[5].text.strip(), "Dr. Christor Pancho")
            expected_date = "September 17, 2026" if u == 1 else "September 24, 2026"
            self.assertEqual(doc.paragraphs[6].text.strip(), expected_date)

    def test_course_folder_processing_fallback_without_metadata(self):
        """Verify fallback when course_metadata.json is absent: infers department from prefix PHIL."""
        course_dir = os.path.join(self.temp_dir, "PHIL1402-Introduction_to_Philosophy")
        os.makedirs(course_dir, exist_ok=True)

        u_folder = os.path.join(course_dir, "01_Unit_1")
        os.makedirs(u_folder, exist_ok=True)
        with open(os.path.join(u_folder, "04_Assignments.md"), "w", encoding="utf-8") as f:
            f.write("# Assignment Activity Unit 1\nEthics and morality paper.")

        result = pac.process_course_folder(course_dir, self.temp_dir)
        self.assertTrue(result)

        assign_dir = os.path.join(u_folder, "coursework", "Assignment_Activity")
        expected_filepath = os.path.join(assign_dir, "week1_assignment_myk_template.docx")
        self.assertTrue(os.path.isfile(expected_filepath))

        import docx
        doc = docx.Document(expected_filepath)
        self.assertEqual(doc.paragraphs[1].text.strip(), "Unit 1 Assignment Activity")
        self.assertEqual(doc.paragraphs[3].text.strip(), "Department of Philosophy, University of The People")
        self.assertIn("PHIL 1402", doc.paragraphs[4].text.strip())

    def test_relative_target_path_population(self):
        """Verify population succeeds when target_path has no directory component (relative filename)."""
        rel_target = f"temp_rel_{int(time.time()*1000)}.docx"
        try:
            success = pac.populate_assignment_template(
                base_template_path=str(self.bundled_template),
                target_path=rel_target,
                title="Unit 3 Assignment Activity",
                department="Department of Computer Science, University of The People",
                course="CS 2401: Software Engineering 1",
                instructor="Dr. Christor Pancho",
                due_date="October 08, 2026"
            )
            self.assertTrue(success)
            self.assertTrue(os.path.isfile(rel_target))
        finally:
            if os.path.exists(rel_target):
                os.remove(rel_target)

    def test_xml_fallback_direct_and_namespaces(self):
        """Verify XML fallback directly populates and preserves docx integrity."""
        target_docx = os.path.join(self.temp_dir, "test_xml_out.docx")
        success = pac.populate_template_xml_fallback(
            base_template_path=str(self.bundled_template),
            target_path=target_docx,
            title="Unit 5 Assignment Activity",
            department="Department of Mathematics, University of The People",
            course="MATH 1201: College Algebra",
            instructor="Prof. Euler",
            due_date="November 12, 2026"
        )
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(target_docx))

        import docx
        doc = docx.Document(target_docx)
        self.assertEqual(doc.paragraphs[1].text.strip(), "Unit 5 Assignment Activity")
        self.assertEqual(doc.paragraphs[3].text.strip(), "Department of Mathematics, University of The People")
        self.assertEqual(doc.paragraphs[4].text.strip(), "MATH 1201: College Algebra")
        self.assertEqual(doc.paragraphs[5].text.strip(), "Prof. Euler")
        self.assertEqual(doc.paragraphs[6].text.strip(), "November 12, 2026")

    def test_semantic_paragraph_fallback_short_doc(self):
        """Verify semantic matching correctly populates documents with <= 6 paragraphs."""
        short_docx = os.path.join(self.temp_dir, "short_template.docx")
        import docx
        base_doc = docx.Document(str(self.bundled_template))
        short_doc = docx.Document()
        # Create minimal 5-paragraph template: Title, Department, Course, Instructor, Date
        short_doc.add_paragraph("x")
        short_doc.add_paragraph("Department of Old Dept, University of The People")
        short_doc.add_paragraph("CS 1101: Programming Fundamentals")
        short_doc.add_paragraph("Instructor Jane Doe")
        short_doc.add_paragraph("July x, 2025")
        short_doc.save(short_docx)

        target_out = os.path.join(self.temp_dir, "short_populated.docx")
        success = pac.populate_assignment_template(
            base_template_path=short_docx,
            target_path=target_out,
            title="Unit 2 Assignment Activity",
            department="Department of Computer Science, University of The People",
            course="CS 2401: Software Engineering 1",
            instructor="Dr. Christor Pancho",
            due_date="October 01, 2026"
        )
        self.assertTrue(success)
        populated = docx.Document(target_out)
        self.assertEqual(populated.paragraphs[0].text.strip(), "Unit 2 Assignment Activity")
        self.assertEqual(populated.paragraphs[1].text.strip(), "Department of Computer Science, University of The People")
        self.assertEqual(populated.paragraphs[2].text.strip(), "CS 2401: Software Engineering 1")
        self.assertEqual(populated.paragraphs[3].text.strip(), "Dr. Christor Pancho")
        self.assertEqual(populated.paragraphs[4].text.strip(), "October 01, 2026")

    def test_unit_regex_boundary_against_unit_10(self):
        r"""Verify regex boundary (?!\d) avoids Unit 1 matching Unit 10, 11, etc."""
        unit_1_regex = re.compile(r'(?:unit|week)\s*0?1(?!\d)', re.I)
        self.assertTrue(unit_1_regex.search("Unit 1 Assignment Activity"))
        self.assertTrue(unit_1_regex.search("Week 01 Assignment"))
        self.assertIsNone(unit_1_regex.search("Unit 10 Assignment Activity"))
        self.assertIsNone(unit_1_regex.search("Unit 11 Assignment Activity"))
        self.assertIsNone(unit_1_regex.search("Week 12 Assignment"))

    def test_student_profile_initials_generation(self):
        """Verify student initials extraction algorithm handles multiple name formats."""
        def extract_initials(full_name):
            parts = [re.sub(r'[^a-zA-Z]', '', p) for p in full_name.split() if re.sub(r'[^a-zA-Z]', '', p)]
            return "".join(p[0].lower() for p in parts) if parts else "myk"

        self.assertEqual(extract_initials("Mohamed Yahia Khidr"), "myk")
        self.assertEqual(extract_initials("John Doe"), "jd")
        self.assertEqual(extract_initials("Jane Alice Doe"), "jad")
        self.assertEqual(extract_initials(""), "myk")


class TestCompanionDaemon(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_port = 4058
        cls.server_address = ("127.0.0.1", cls.test_port)
        cls.httpd = daemon.ThreadingHTTPServer(cls.server_address, daemon.CourseExporterHandler)
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()
        time.sleep(0.5)

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.thread.join(timeout=2.0)
        time.sleep(0.2)

    def test_health_endpoint(self):
        """Verify GET /health returns service status and port."""
        req = urllib.request.Request(f"http://127.0.0.1:{self.test_port}/health")
        with urllib.request.urlopen(req, timeout=3) as resp:
            self.assertEqual(resp.status, 200)
            cors = resp.headers.get("Access-Control-Allow-Origin")
            self.assertEqual(cors, "*")
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(data.get("status"), "ok")
            self.assertEqual(data.get("service"), "uopeople-course-exporter-daemon")

    def test_cors_options(self):
        """Verify OPTIONS preflight returns CORS headers."""
        req = urllib.request.Request(f"http://127.0.0.1:{self.test_port}/process-courses", method="OPTIONS")
        with urllib.request.urlopen(req, timeout=3) as resp:
            self.assertIn(resp.status, (200, 204))
            self.assertEqual(resp.headers.get("Access-Control-Allow-Origin"), "*")
            self.assertIn("POST", resp.headers.get("Access-Control-Allow-Methods", ""))

    def test_process_courses_trigger(self):
        """Verify POST /process-courses responds immediately with triggered status."""
        payload = {
            "fileName": "UoPeople_CS_2401_Offline.zip",
            "courseName": "CS 2401 Software Engineering 1",
            "activeCoursesFolder": ""
        }
        req = urllib.request.Request(
            f"http://127.0.0.1:{self.test_port}/process-courses",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST"
        )
        with urllib.request.urlopen(req, timeout=3) as resp:
            self.assertEqual(resp.status, 200)
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(data.get("status"), "triggered")


if __name__ == "__main__":
    unittest.main()
